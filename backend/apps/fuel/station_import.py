"""Admin spreadsheet import for fuel stations: upload → validated preview → confirmed apply.

The command-line importer (``apps.fuel.importer``) loads the full employer CSV. This module is
the administrator-facing path for the same data and shares its column names, value validation
(OPIS IDs, rack IDs, US state codes, exact Decimal prices) and duplicate policy. Differences
that exist because an administrator may upload partial files:

* Only ``OPIS Truckstop ID`` is a required column. Omitted columns, and blank cells, never
  clear stored values: they mean "keep what is there".
* A new station (OPIS ID not in the database) needs Name, City, State, Rack ID and Retail
  Price; otherwise its rows are rejected.
* Rows sharing an OPIS ID are merged exactly like the CLI (lowest price, highest price and
  observation count; most frequent name); rows that disagree on location are all rejected.
* Coordinates are never read from files. New or relocated stations are geolocated offline
  from the Census gazetteer (town-level), exactly as after a CLI import.

Invalid rows never block the valid ones: they are listed in the preview and the error report,
and the valid rows are applied together in one transaction (all or nothing).
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import logging
import math
import time
import uuid
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass, field
from decimal import Decimal
from pathlib import PurePath
from typing import Any

from django.conf import settings
from django.db import transaction
from django.utils import timezone

from apps.fuel.constants import CANADIAN_PROVINCE_CODES, US_STATE_CODES
from apps.fuel.geolocation import Gazetteer, GazetteerError, geolocate_stations
from apps.fuel.importer import _HAS_ALNUM as HAS_ALNUM
from apps.fuel.importer import (
    BULK_BATCH_SIZE,
    COL_ADDRESS,
    COL_CITY,
    COL_ID,
    COL_NAME,
    COL_PRICE,
    COL_RACK,
    COL_STATE,
    REQUIRED_COLUMNS,
    Reason,
    RowIssue,
    normalize_text,
    parse_price,
)
from apps.fuel.importer import _parse_positive_int as parse_positive_int
from apps.fuel.importer import _reset_geolocation as reset_geolocation
from apps.fuel.models import FuelStation, StationImport

logger = logging.getLogger(__name__)

# --- limits ---------------------------------------------------------------------------------

MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # the full employer CSV is ~0.6 MB
MAX_DATA_ROWS = 20_000  # the employer CSV has 8,151
MAX_COLUMNS = 30
MAX_XLSX_ENTRIES = 500
MAX_XLSX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024  # zip-bomb guard
MAX_SHEET_ROWS_SCANNED = 200_000  # includes blank rows
PREVIEW_TTL = dt.timedelta(minutes=30)
PREVIEW_RETENTION = dt.timedelta(days=1)
SAMPLE_SIZE = 50  # representative changes shown in the preview
TOO_LARGE_MESSAGE = f"The file is larger than the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit."

ALLOWED_SUFFIXES = (".csv", ".xlsx")
_MACRO_SUFFIXES = (".xlsm", ".xltm", ".xlam", ".xlsb")
_ZIP_MAGIC = b"PK\x03\x04"
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # legacy .xls and encrypted workbooks
XLSX_SHEET_NAME = "Stations"

# --- columns ---------------------------------------------------------------------------------

TEMPLATE_COLUMNS = REQUIRED_COLUMNS  # same headers, same order as the employer CSV
FIELD_BY_COLUMN = {
    COL_ID: "opis_id",
    COL_NAME: "name",
    COL_ADDRESS: "address",
    COL_CITY: "city",
    COL_STATE: "state",
    COL_RACK: "rack_id",
    COL_PRICE: "price",
}
NEW_STATION_FIELDS = ("name", "city", "state", "rack_id", "price")
LOCATION_FIELDS = ("address", "city", "state")
FIELD_LABELS = {
    "name": "Truckstop Name",
    "address": "Address",
    "city": "City",
    "state": "State",
    "rack_id": "Rack ID",
    "price": "Retail Price",
    "retail_price": "Retail Price",
    "max_retail_price": "Highest listed price",
    "price_observations": "Price observations",
}
PRICE_FIELDS = ("retail_price", "max_retail_price")
_TEXT_LIMITS = {
    name: FuelStation._meta.get_field(name).max_length for name in ("name", "address", "city")
}
_COORDINATE_HINTS = ("lat", "lon", "lng", "coordinate", "geolocation")

REASON_LABELS = {
    Reason.MALFORMED_ROW: "Malformed row",
    Reason.MISSING_FIELD: "Missing value",
    Reason.INVALID_ID: "Invalid OPIS ID",
    Reason.INVALID_RACK_ID: "Invalid Rack ID",
    Reason.INVALID_STATE: "Invalid state",
    Reason.INVALID_PRICE: "Invalid price",
    Reason.CONFLICTING_DUPLICATE: "Conflicting duplicate",
    Reason.NON_US_STATION: "Outside the US",
    Reason.UNSUPPORTED_VALUE: "Unsupported cell",
    Reason.VALUE_TOO_LONG: "Value too long",
    Reason.INCOMPLETE_NEW_STATION: "Incomplete new station",
}


class UploadError(Exception):
    """The file cannot be imported at all. The message is safe to show to the user."""


class ImportRejected(Exception):
    """A confirmation was refused; nothing was written. The message is safe to show."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


# --- reading files ---------------------------------------------------------------------------


@dataclass(frozen=True)
class CellProblem:
    """A spreadsheet cell whose value cannot be imported (formula, date, error, boolean)."""

    message: str


Cell = str | CellProblem


@dataclass
class Sheet:
    file_format: str
    header: list[str]
    rows: list[tuple[int, list[Cell]]]  # (line or row number, cells); blank rows removed
    strict_width: bool  # CSV rows must have exactly as many cells as the header


def read_upload(file_name: str, data: bytes) -> Sheet:
    """Detect the real format from the content, enforce limits and return the cells."""
    suffix = PurePath(file_name).suffix.lower()
    if suffix in _MACRO_SUFFIXES:
        raise UploadError(
            "Macro-enabled or binary workbooks are not accepted. "
            "Save the file as an Excel Workbook (.xlsx) or CSV UTF-8 and upload it again."
        )
    if suffix == ".xls":
        raise UploadError("Legacy .xls workbooks are not supported. Save as .xlsx or CSV UTF-8.")
    if suffix not in ALLOWED_SUFFIXES:
        raise UploadError("Unsupported file type. Upload a .csv or .xlsx file.")
    if not data:
        raise UploadError("The file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise UploadError(TOO_LARGE_MESSAGE)
    if suffix == ".xlsx":
        if not data.startswith(_ZIP_MAGIC):
            raise UploadError(
                "The file is named .xlsx but is not an Excel workbook "
                "(password-protected workbooks are not supported)."
            )
        return _read_xlsx(data)
    if data.startswith((_ZIP_MAGIC, _OLE_MAGIC)) or b"\x00" in data:
        raise UploadError(
            "The file is named .csv but contains binary data. "
            "Upload the workbook as .xlsx, or save it as CSV UTF-8."
        )
    return _read_csv(data)


def _read_csv(data: bytes) -> Sheet:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise UploadError(
            "The CSV is not UTF-8 encoded. In Excel use Save As → “CSV UTF-8 (Comma delimited)”."
        ) from exc
    reader = csv.reader(io.StringIO(text, newline=""))
    header: list[str] | None = None
    rows: list[tuple[int, list[Cell]]] = []
    try:
        for cells in reader:
            if header is None:
                header = cells
                continue
            if not any(cell.strip() for cell in cells):
                continue
            if len(rows) >= MAX_DATA_ROWS:
                raise UploadError(f"The file has more than {MAX_DATA_ROWS:,} data rows.")
            rows.append((reader.line_num, list(cells)))
    except csv.Error as exc:
        raise UploadError(f"The CSV could not be parsed (line {reader.line_num}): {exc}.") from exc
    if header is None:
        raise UploadError("The file is empty.")
    return Sheet("csv", header, rows, strict_width=True)


def _read_xlsx(data: bytes) -> Sheet:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise UploadError("The workbook is damaged or is not a valid .xlsx file.") from exc
    with archive:
        entries = archive.infolist()
        names = {entry.filename for entry in entries}
        if len(entries) > MAX_XLSX_ENTRIES or (
            sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES
        ):
            raise UploadError("The workbook expands to more data than the import allows.")
        if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
            raise UploadError("The file is not a valid .xlsx workbook.")
        if any(name.lower().endswith("vbaproject.bin") for name in names) or (
            b"macroenabled" in archive.read("[Content_Types].xml").lower()
        ):
            raise UploadError(
                "Macro-enabled workbooks are not accepted. Save as an Excel Workbook (.xlsx)."
            )
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on the deployment
        raise UploadError(
            "Excel support (openpyxl) is not installed on this server. Upload a CSV file."
        ) from exc

    try:
        # read_only streams rows; data_only=False exposes formulas so they can be refused
        # (openpyxl never evaluates formulas). External links are not loaded.
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many exception types
        logger.info("Rejected unreadable workbook: %s", type(exc).__name__)
        raise UploadError("The workbook could not be read. Re-save it as .xlsx and retry.") from exc
    try:
        if XLSX_SHEET_NAME in workbook.sheetnames:
            worksheet = workbook[XLSX_SHEET_NAME]
        elif workbook.worksheets:
            worksheet = workbook.worksheets[0]
        else:
            raise UploadError("The workbook has no worksheets.")
        worksheet.reset_dimensions()  # never trust the declared used range
        header: list[str] | None = None
        rows: list[tuple[int, list[Cell]]] = []
        scanned = 0
        for row in worksheet.iter_rows():
            scanned += 1
            if scanned > MAX_SHEET_ROWS_SCANNED:
                raise UploadError("The worksheet has too many rows.")
            if len(row) > MAX_COLUMNS * 4:
                raise UploadError(f"The worksheet has more than {MAX_COLUMNS} columns.")
            cells = [_xlsx_cell(cell) for cell in row]
            if not any(isinstance(c, CellProblem) or c.strip() for c in cells):
                continue
            number = next((c.row for c in row if getattr(c, "row", None)), scanned)
            if header is None:
                if any(isinstance(c, CellProblem) for c in cells):
                    raise UploadError("The header row must contain plain text column names.")
                header = [c for c in cells if isinstance(c, str)]
                continue
            if len(rows) >= MAX_DATA_ROWS:
                raise UploadError(f"The file has more than {MAX_DATA_ROWS:,} data rows.")
            rows.append((number, cells))
    except UploadError:
        raise
    except Exception as exc:  # noqa: BLE001 - malformed sheet XML surfaces while streaming
        logger.info("Rejected malformed worksheet: %s", type(exc).__name__)
        raise UploadError(
            "The worksheet could not be read. Re-save it as .xlsx and retry."
        ) from exc
    finally:
        workbook.close()
    if header is None:
        raise UploadError(f"The worksheet “{worksheet.title}” is empty.")
    while header and not header[-1].strip():
        header.pop()
    return Sheet("xlsx", header, rows, strict_width=False)


def _xlsx_cell(cell: Any) -> Cell:
    value = cell.value
    if value is None:
        return ""
    data_type = getattr(cell, "data_type", None)
    if data_type == "f":
        return CellProblem("contains a formula; paste values only")
    if data_type == "e":
        return CellProblem(f"contains the error value {value}")
    if isinstance(value, bool):
        return CellProblem("contains TRUE/FALSE")
    if isinstance(value, dt.datetime | dt.date | dt.time | dt.timedelta):
        return CellProblem("contains a date or time")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return CellProblem("is not a finite number")
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)  # shortest round-trip text, e.g. 3.459 → "3.459"
    if isinstance(value, str):
        return value
    return CellProblem("has an unsupported value type")


# --- header and row validation -------------------------------------------------------------


@dataclass
class ColumnMap:
    positions: dict[str, int]  # field → column index
    ignored: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    @property
    def columns(self) -> list[str]:
        by_field = {v: k for k, v in FIELD_BY_COLUMN.items()}
        return [by_field[f] for f in sorted(self.positions, key=self.positions.__getitem__)]


def resolve_columns(header: list[str]) -> ColumnMap:
    if len(header) > MAX_COLUMNS:
        raise UploadError(f"The file has more than {MAX_COLUMNS} columns.")
    lookup = {column.casefold(): column for column in FIELD_BY_COLUMN}
    columns = ColumnMap(positions={})
    for index, raw in enumerate(header):
        text = normalize_text(raw)
        if not text:
            continue
        canonical = lookup.get(text.casefold())
        if canonical is None:
            columns.ignored.append(text)
            continue
        name = FIELD_BY_COLUMN[canonical]
        if name in columns.positions:
            raise UploadError(f"The column “{canonical}” appears more than once.")
        columns.positions[name] = index
    if "opis_id" not in columns.positions:
        hint = ""
        if any(";" in cell for cell in header):
            hint = " The file looks semicolon-separated; save it as CSV UTF-8 (comma-separated)."
        raise UploadError(f"The required column “{COL_ID}” was not found.{hint}")
    if len(columns.positions) == 1:
        raise UploadError(
            "The file only has the OPIS Truckstop ID column; add the columns to import."
        )
    if columns.ignored:
        columns.notes.append(f"Ignored unknown columns: {', '.join(columns.ignored)}.")
        if any(any(h in c.casefold() for h in _COORDINATE_HINTS) for c in columns.ignored):
            columns.notes.append(
                "Coordinates are never imported from files. New or relocated stations are "
                "placed at their town from the offline Census gazetteer."
            )
    return columns


@dataclass(frozen=True)
class IncomingRow:
    line: int
    opis_id: int
    name: str | None
    address: str | None
    city: str | None
    state: str | None
    rack_id: int | None
    price: Decimal | None


def validate_row(
    line: int, cells: list[Cell], columns: ColumnMap, sheet: Sheet
) -> IncomingRow | RowIssue:
    """Validate one data row. ``None`` values mean "not provided" (column absent or blank)."""
    width = len(sheet.header)
    id_cell = _cell(cells, columns.positions["opis_id"])
    id_text = normalize_text(id_cell) if isinstance(id_cell, str) else ""

    if sheet.strict_width and len(cells) != width:
        return RowIssue(
            line, id_text, Reason.MALFORMED_ROW, f"expected {width} cells, found {len(cells)}"
        )
    if any(isinstance(c, CellProblem) or c.strip() for c in cells[width:]):
        return RowIssue(line, id_text, Reason.MALFORMED_ROW, "values outside the header columns")

    values: dict[str, str | None] = {}
    for name, index in columns.positions.items():
        cell = _cell(cells, index)
        if isinstance(cell, CellProblem):
            label = FIELD_LABELS.get(name, COL_ID)
            return RowIssue(line, id_text, Reason.UNSUPPORTED_VALUE, f"{label} {cell.message}")
        values[name] = normalize_text(cell) or None

    if values["opis_id"] is None:
        return RowIssue(line, id_text, Reason.MISSING_FIELD, f"{COL_ID} is required")
    opis_id = parse_positive_int(values["opis_id"])
    if opis_id is None:
        return RowIssue(line, id_text, Reason.INVALID_ID, f"not a positive integer: {id_text!r}")

    state = values.get("state")
    if state is not None:
        state = state.upper()
        if state in CANADIAN_PROVINCE_CODES:
            return RowIssue(line, id_text, Reason.NON_US_STATION, f"Canadian province {state}")
        if state not in US_STATE_CODES:
            return RowIssue(line, id_text, Reason.INVALID_STATE, f"unknown state code {state!r}")

    rack_id = None
    if values.get("rack_id") is not None:
        rack_id = parse_positive_int(values["rack_id"])
        if rack_id is None:
            return RowIssue(
                line,
                id_text,
                Reason.INVALID_RACK_ID,
                f"not a positive integer: {values['rack_id']!r}",
            )

    price = None
    if values.get("price") is not None:
        try:
            price = parse_price(values["price"])
        except ValueError as exc:
            return RowIssue(line, id_text, Reason.INVALID_PRICE, str(exc))

    for name, limit in _TEXT_LIMITS.items():
        text = values.get(name)
        if text is not None and len(text) > limit:
            return RowIssue(
                line, id_text, Reason.VALUE_TOO_LONG, f"{FIELD_LABELS[name]} exceeds {limit} chars"
            )

    address = values.get("address")
    if address is not None and not HAS_ALNUM.search(address):
        address = None  # "-" or "N/A"-style placeholders are treated as blank
    return IncomingRow(
        line=line,
        opis_id=opis_id,
        name=values.get("name"),
        address=address,
        city=values.get("city"),
        state=state,
        rack_id=rack_id,
        price=price,
    )


def _cell(cells: list[Cell], index: int) -> Cell:
    return cells[index] if index < len(cells) else ""


# --- duplicate merging (same policy as the CLI importer) ------------------------------------


@dataclass(frozen=True)
class StationRecord:
    """Everything the file says about one station. ``None`` fields keep stored values."""

    opis_id: int
    lines: tuple[int, ...]
    name: str | None
    address: str | None
    city: str | None
    state: str | None
    rack_id: int | None
    retail_price: Decimal | None
    max_retail_price: Decimal | None
    price_observations: int | None

    def to_json(self) -> dict[str, Any]:
        data = asdict(self)
        data["lines"] = list(self.lines)
        for name in PRICE_FIELDS:
            data[name] = None if data[name] is None else str(data[name])
        return data

    @classmethod
    def from_json(cls, data: dict[str, Any]) -> StationRecord:
        values = dict(data)
        values["lines"] = tuple(values["lines"])
        for name in PRICE_FIELDS:
            values[name] = None if values[name] is None else Decimal(values[name])
        return cls(**values)


def _location_key(row: IncomingRow) -> tuple[object, ...]:
    def fold(text: str | None) -> str | None:
        return None if text is None else text.casefold()

    return (fold(row.address), fold(row.city), row.state, row.rack_id)


def merge_rows(rows: list[IncomingRow]) -> tuple[list[StationRecord], list[RowIssue], int]:
    """Collapse rows per OPIS ID. Returns (records, conflict issues, rows merged away)."""
    groups: dict[int, list[IncomingRow]] = {}
    for row in rows:
        groups.setdefault(row.opis_id, []).append(row)
    records: list[StationRecord] = []
    conflicts: list[RowIssue] = []
    merged = 0
    for opis_id, group in groups.items():
        lines = tuple(row.line for row in group)
        if len({_location_key(row) for row in group}) > 1:
            listed = ", ".join(str(line) for line in lines)
            conflicts.extend(
                RowIssue(
                    row.line,
                    str(opis_id),
                    Reason.CONFLICTING_DUPLICATE,
                    f"location or rack differs across rows {listed}",
                )
                for row in group
            )
            continue
        merged += len(group) - 1
        first = group[0]
        names = [row.name for row in group if row.name is not None]
        prices = [row.price for row in group if row.price is not None]
        records.append(
            StationRecord(
                opis_id=opis_id,
                lines=lines,
                name=Counter(names).most_common(1)[0][0] if names else None,
                address=first.address,
                city=first.city,
                state=first.state,
                rack_id=first.rack_id,
                retail_price=min(prices) if prices else None,
                max_retail_price=max(prices) if prices else None,
                price_observations=len(prices) if prices else None,
            )
        )
    return records, conflicts, merged


# --- planning --------------------------------------------------------------------------------


@dataclass
class Plan:
    creates: list[dict[str, Any]] = field(default_factory=list)
    updates: list[dict[str, Any]] = field(default_factory=list)
    unchanged: list[int] = field(default_factory=list)
    blocked: list[RowIssue] = field(default_factory=list)
    snapshot: dict[str, str] = field(default_factory=dict)  # opis_id → stored updated_at
    stations: dict[int, FuelStation] = field(default_factory=dict, repr=False)  # by opis_id

    @property
    def fingerprint(self) -> str:
        """Hash of every change and of the stored versions it was computed against."""
        payload = {
            "creates": self.creates,
            "updates": [
                {k: u[k] for k in ("opis_id", "changes", "relocate")} for u in self.updates
            ],
            "unchanged": self.unchanged,
            "snapshot": self.snapshot,
        }
        text = json.dumps(payload, sort_keys=True, separators=(",", ":"), default=str)
        return hashlib.sha256(text.encode()).hexdigest()


def _json_value(value: object) -> object:
    return str(value) if isinstance(value, Decimal) else value


def _same(stored: object, incoming: object) -> bool:
    if isinstance(stored, str) and isinstance(incoming, str):
        return normalize_text(stored) == incoming
    return stored == incoming


def build_plan(records: list[StationRecord], *, lock: bool = False) -> Plan:
    """Compare records with the database (one batched lookup) and decide each station's fate."""
    queryset = FuelStation.objects.select_for_update() if lock else FuelStation.objects.all()
    existing = queryset.in_bulk([r.opis_id for r in records], field_name="opis_id")
    plan = Plan()
    for record in records:
        station = existing.get(record.opis_id)
        if station is None:
            missing = [
                FIELD_LABELS[f] for f in NEW_STATION_FIELDS if _record_value(record, f) is None
            ]
            if missing:
                detail = f"new station (OPIS ID not in the database) needs {', '.join(missing)}"
                plan.blocked.extend(
                    RowIssue(line, str(record.opis_id), Reason.INCOMPLETE_NEW_STATION, detail)
                    for line in record.lines
                )
                continue
            plan.creates.append(
                {
                    "opis_id": record.opis_id,
                    "name": record.name,
                    "address": record.address or "",
                    "city": record.city,
                    "state": record.state,
                    "rack_id": record.rack_id,
                    "retail_price": str(record.retail_price),
                    "max_retail_price": str(record.max_retail_price),
                    "price_observations": record.price_observations,
                }
            )
            continue

        plan.stations[record.opis_id] = station
        plan.snapshot[str(record.opis_id)] = station.updated_at.isoformat()
        changes = []
        candidates = ["name", "address", "city", "state", "rack_id"]
        if record.retail_price is not None:
            candidates += ["retail_price", "max_retail_price", "price_observations"]
        for name in candidates:
            incoming = getattr(record, name)
            stored = getattr(station, name)
            if incoming is None or _same(stored, incoming):
                continue
            changes.append(
                {"field": name, "old": _json_value(stored), "new": _json_value(incoming)}
            )
        if not changes:
            plan.unchanged.append(record.opis_id)
            continue
        plan.updates.append(
            {
                "opis_id": record.opis_id,
                "name": station.name,
                "changes": changes,
                # Old coordinates belong to the old location; never keep them after a move.
                "relocate": any(c["field"] in LOCATION_FIELDS for c in changes),
                "had_coordinates": station.has_coordinates,
            }
        )
    return plan


def _record_value(record: StationRecord, name: str) -> object:
    return record.retail_price if name == "price" else getattr(record, name)


# --- preview ---------------------------------------------------------------------------------


def _issue_json(issue: RowIssue) -> dict[str, Any]:
    return {
        "line": issue.line,
        "opis_id": issue.opis_id,
        "reason": issue.reason.value,
        "label": REASON_LABELS.get(issue.reason, issue.reason.value),
        "detail": issue.detail,
    }


def create_preview(user: Any, file_name: str, data: bytes) -> StationImport:
    """Validate an upload and store its preview. Writes nothing to ``FuelStation``.

    Raises ``UploadError`` when the file as a whole cannot be imported.
    """
    t0 = time.perf_counter()
    sheet = read_upload(file_name, data)
    t1 = time.perf_counter()
    columns = resolve_columns(sheet.header)
    rows: list[IncomingRow] = []
    issues: list[RowIssue] = []
    for line, cells in sheet.rows:
        outcome = validate_row(line, cells, columns, sheet)
        if isinstance(outcome, IncomingRow):
            rows.append(outcome)
        else:
            issues.append(outcome)
    records, conflicts, merged = merge_rows(rows)
    t2 = time.perf_counter()
    plan = build_plan(records)
    t3 = time.perf_counter()

    issues = sorted([*issues, *conflicts, *plan.blocked], key=lambda issue: issue.line)
    summary = {
        "total_rows": len(sheet.rows),
        "columns": columns.columns,
        "notes": columns.notes,
        "stations_in_file": len(records),
        "new": len(plan.creates),
        "updates": len(plan.updates),
        "unchanged": len(plan.unchanged),
        "invalid_rows": len(issues),
        "invalid_by_reason": dict(Counter(issue.reason.value for issue in issues).most_common()),
        "duplicate_rows": merged + len(conflicts),
        "duplicate_rows_merged": merged,
        "conflicting_duplicate_rows": len(conflicts),
        "relocated": sum(1 for update in plan.updates if update["relocate"]),
        "changed_fields": dict(
            Counter(c["field"] for u in plan.updates for c in u["changes"]).most_common()
        ),
        "sample_updates": plan.updates[:SAMPLE_SIZE],
        "sample_creates": plan.creates[:SAMPLE_SIZE],
        "needs_add": bool(plan.creates),
        "needs_change": bool(plan.updates),
        "timings_ms": {
            "parse": _ms(t1 - t0),
            "validate": _ms(t2 - t1),
            "match": _ms(t3 - t2),
        },
    }
    with transaction.atomic():
        StationImport.objects.filter(
            status=StationImport.Status.PREVIEWED, created_at__lt=timezone.now() - PREVIEW_RETENTION
        ).delete()
        return StationImport.objects.create(
            created_by=user,
            file_name=PurePath(file_name).name[:255],
            file_format=sheet.file_format,
            file_size=len(data),
            file_sha256=hashlib.sha256(data).hexdigest(),
            fingerprint=plan.fingerprint,
            rows=[record.to_json() for record in records],
            summary=summary,
            issues=[_issue_json(issue) for issue in issues],
        )


def _ms(seconds: float) -> float:
    return round(seconds * 1000, 1)


# --- confirmation ----------------------------------------------------------------------------


def missing_permissions(user: Any, summary: dict[str, Any]) -> list[str]:
    """Human-readable permissions the user lacks for the changes in a preview."""
    missing = []
    if summary.get("needs_add") and not user.has_perm("fuel.add_fuelstation"):
        missing.append("add fuel stations")
    if summary.get("needs_change") and not user.has_perm("fuel.change_fuelstation"):
        missing.append("change fuel stations")
    return missing


def load_gazetteer() -> Gazetteer | None:
    try:
        return Gazetteer.from_file(settings.GAZETTEER_PLACES_PATH)
    except GazetteerError:
        logger.warning("Gazetteer unavailable; imported stations stay unresolved.", exc_info=True)
        return None


class _StalePreview(Exception):
    pass


def confirm_import(import_id: uuid.UUID, user: Any, fingerprint: str) -> StationImport:
    """Apply a preview exactly as shown, atomically. Raises ``ImportRejected`` otherwise.

    * Only the uploader can confirm, once (the status claim is the double-submit guard).
    * The plan is recomputed inside the transaction (rows locked where the database supports
      it); if it differs from the preview in any way, nothing is written.
    """
    station_import = StationImport.objects.filter(pk=import_id, created_by=user).first()
    if station_import is None:
        raise ImportRejected("not_found", "This import preview does not exist.")
    if station_import.status != StationImport.Status.PREVIEWED:
        raise ImportRejected(
            "already_processed", "This preview was already confirmed or closed. Nothing changed."
        )
    if fingerprint != station_import.fingerprint:
        raise ImportRejected(
            "mismatch", "The confirmation does not match this preview. Reload it and try again."
        )
    now = timezone.now()
    if now - station_import.created_at > PREVIEW_TTL:
        _close(station_import, StationImport.Status.EXPIRED)
        raise ImportRejected(
            "expired", "This preview expired. Upload the file again to see current changes."
        )
    missing = missing_permissions(user, station_import.summary)
    if missing:
        raise ImportRejected("permission", f"You do not have permission to {' or '.join(missing)}.")

    records = [StationRecord.from_json(data) for data in station_import.rows]
    summary = station_import.summary
    needs_geolocation = summary.get("new", 0) or summary.get("relocated", 0)
    gazetteer = load_gazetteer() if needs_geolocation else None  # loaded outside the transaction
    started = time.perf_counter()
    try:
        with transaction.atomic():
            # Claim first: this UPDATE also takes SQLite's write lock, so no other writer can
            # commit between the re-check below and our writes.
            claimed = StationImport.objects.filter(
                pk=station_import.pk, status=StationImport.Status.PREVIEWED
            ).update(status=StationImport.Status.APPLIED, finished_at=now)
            if claimed != 1:
                raise ImportRejected(
                    "already_processed",
                    "This preview was already confirmed or closed. Nothing changed.",
                )
            plan = build_plan(records, lock=True)
            if plan.fingerprint != station_import.fingerprint:
                raise _StalePreview
            inserted, updated = _write(plan)
            geo = None
            touched = [c["opis_id"] for c in plan.creates] + [
                u["opis_id"] for u in plan.updates if u["relocate"]
            ]
            if gazetteer is not None and touched:
                geo = geolocate_stations(gazetteer, opis_ids=touched)
    except _StalePreview:
        _close(station_import, StationImport.Status.STALE)
        raise ImportRejected(
            "stale",
            "Station data changed after this preview was created (an edit or another import). "
            "Nothing was imported. Upload the file again to review the current changes.",
        ) from None

    station_import.refresh_from_db()
    station_import.result = {
        "total_rows": summary["total_rows"],
        "inserted": inserted,
        "updated": updated,
        "unchanged": len(plan.unchanged),
        "rejected": summary["invalid_rows"],
        "duplicate_rows": summary["duplicate_rows"],
        "relocated": summary["relocated"],
        "geolocation": "skipped" if needs_geolocation and gazetteer is None else "applied",
        "geolocated": geo.resolved if geo else 0,
        "left_unresolved": (geo.considered - geo.resolved) if geo else len(touched),
        "duration_ms": _ms(time.perf_counter() - started),
    }
    station_import.rows = []  # the applied data now lives in FuelStation
    station_import.save(update_fields=["result", "rows"])
    return station_import


def _write(plan: Plan) -> tuple[int, int]:
    now = timezone.now()
    created = [
        FuelStation(
            **{
                **values,
                "retail_price": Decimal(values["retail_price"]),
                "max_retail_price": Decimal(values["max_retail_price"]),
            }
        )
        for values in plan.creates
    ]
    FuelStation.objects.bulk_create(created, batch_size=BULK_BATCH_SIZE)

    changed_fields: set[str] = {"updated_at"}
    to_update = []
    for update in plan.updates:
        station = plan.stations[update["opis_id"]]
        for change in update["changes"]:
            value = change["new"]
            setattr(
                station,
                change["field"],
                Decimal(value) if change["field"] in PRICE_FIELDS else value,
            )
            changed_fields.add(change["field"])
        if update["relocate"]:
            reset_geolocation(station)
            changed_fields.update(
                (
                    "latitude",
                    "longitude",
                    "geolocation_source",
                    "geolocation_accuracy",
                    "location_uncertainty_miles",
                )
            )
        station.updated_at = now  # bumps the station-data version the route planner checks
        to_update.append(station)
    FuelStation.objects.bulk_update(to_update, sorted(changed_fields), batch_size=BULK_BATCH_SIZE)
    return len(created), len(to_update)


def _close(station_import: StationImport, status: str) -> None:
    StationImport.objects.filter(
        pk=station_import.pk, status=StationImport.Status.PREVIEWED
    ).update(status=status, finished_at=timezone.now(), rows=[])
    station_import.refresh_from_db()


# --- downloads -------------------------------------------------------------------------------

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe_csv_cell(value: object) -> str:
    """Neutralize spreadsheet formula injection in exported text."""
    text = "" if value is None else str(value)
    return f"'{text}" if text.startswith(_FORMULA_PREFIXES) else text


def error_report_csv(station_import: StationImport) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["Row", COL_ID, "Reason", "Code", "Detail"])
    for issue in station_import.issues:
        writer.writerow(
            [
                issue["line"],
                _safe_csv_cell(issue["opis_id"]),
                issue["label"],
                issue["reason"],
                _safe_csv_cell(issue["detail"]),
            ]
        )
    return buffer.getvalue()


def csv_template() -> str:
    buffer = io.StringIO()
    csv.writer(buffer).writerow(TEMPLATE_COLUMNS)
    return buffer.getvalue()


FIELD_GUIDE = (
    (COL_ID, "Required", "Required", "Positive whole number. Matches the existing station."),
    (COL_NAME, "Required", "Optional", "Up to 255 characters."),
    (COL_ADDRESS, "Optional", "Optional", "Up to 255 characters. A change re-derives coordinates."),
    (COL_CITY, "Required", "Optional", "Up to 100 characters. A change re-derives coordinates."),
    (COL_STATE, "Required", "Optional", "Two-letter US state code (DC allowed)."),
    (COL_RACK, "Required", "Optional", "Positive whole number."),
    (COL_PRICE, "Required", "Optional", "USD per gallon, 0.50–20.00, up to 8 decimal places."),
)
EXAMPLE_ROWS = (
    (999_999_901, "EXAMPLE TRAVEL CENTER", "I-35, EXIT 12", "Dallas", "TX", 101, 3.459),
    (999_999_902, "EXAMPLE TRUCK STOP", "I-20, EXIT 4", "Abilene", "TX", 102, 3.299),
)


def xlsx_template() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="397C70")
    workbook = Workbook()
    stations = workbook.active
    stations.title = XLSX_SHEET_NAME
    stations.append(list(TEMPLATE_COLUMNS))
    for cell in stations[1]:
        cell.font, cell.fill = header_font, header_fill
    for letter, width in zip("ABCDEFG", (20, 34, 30, 20, 8, 10, 14), strict=True):
        stations.column_dimensions[letter].width = width
    stations.freeze_panes = "A2"

    example = workbook.create_sheet("Example")
    example.append(
        ["EXAMPLE ROWS ONLY — this sheet is never imported. Fill in the Stations sheet."]
    )
    example["A1"].font = Font(bold=True, color="84591E")
    example.append(list(TEMPLATE_COLUMNS))
    for cell in example[2]:
        cell.font = Font(bold=True)
    for row in EXAMPLE_ROWS:
        example.append(list(row))
    for letter, width in zip("ABCDEFG", (20, 34, 30, 20, 8, 10, 14), strict=True):
        example.column_dimensions[letter].width = width

    guide = workbook.create_sheet("Field guide")
    guide.append(["Column", "New station", "Existing station", "Format and rules"])
    for cell in guide[1]:
        cell.font, cell.fill = header_font, header_fill
    for row in FIELD_GUIDE:
        guide.append(list(row))
    notes = (
        "Only the Stations sheet is imported. Keep the header row exactly as provided.",
        "Existing stations are matched by OPIS Truckstop ID only.",
        "Blank cells and omitted columns keep the stored value; nothing is cleared.",
        "Rows repeating an OPIS ID are merged (lowest/highest price) if their location matches.",
        "Coordinates are not imported; stations are placed at their town automatically.",
        "Enter values, not formulas. Dates and TRUE/FALSE cells are rejected.",
    )
    guide.append([])
    for note in notes:
        guide.append([note])
    for letter, width in zip("ABCD", (20, 14, 16, 70), strict=True):
        guide.column_dimensions[letter].width = width
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
