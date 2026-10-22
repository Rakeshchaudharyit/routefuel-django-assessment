"""Admin CSV/XLSX fuel-station import: security, validation, preview, confirmation, integrity."""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import timedelta
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import connection
from django.test import Client
from django.test.utils import CaptureQueriesContext

from apps.fuel import station_import
from apps.fuel.importer import REQUIRED_COLUMNS
from apps.fuel.models import FuelStation, GeolocationAccuracy, GeolocationSource, StationImport
from apps.fuel.queries import routable_stations
from apps.routing.fuel_plans import _station_data_version

pytestmark = pytest.mark.django_db

FIXTURES = Path(__file__).parent / "fixtures"
IMPORT_URL = "/admin/fuel/fuelstation/import/"
HEADER = list(REQUIRED_COLUMNS)


@pytest.fixture(autouse=True)
def sample_gazetteer_setting(settings):
    settings.GAZETTEER_PLACES_PATH = FIXTURES / "gazetteer_sample.txt"


# --- helpers ---------------------------------------------------------------------------------


def make_user(django_user_model, *codenames, name="importer"):
    # Sessions come from force_login, so no password is set (and none is hashed).
    user = django_user_model.objects.create_user(name, password=None, is_staff=True)
    user.user_permissions.add(
        *Permission.objects.filter(content_type__app_label="fuel", codename__in=codenames)
    )
    return user, None


@pytest.fixture
def admin_user(django_user_model):
    return make_user(django_user_model, "view_fuelstation", "add_fuelstation", "change_fuelstation")


@pytest.fixture
def admin_client(admin_user):
    client = Client()
    client.force_login(admin_user[0])
    return client


def csv_bytes(rows, header=HEADER) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode()


def xlsx_bytes(rows, header=HEADER, sheet="Stations", extra_sheets=()) -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.append(header)
    for row in rows:
        worksheet.append(row)
    for name, sheet_rows in extra_sheets:
        other = workbook.create_sheet(name)
        for row in sheet_rows:
            other.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def upload(client, data: bytes, name="stations.csv"):
    return client.post(IMPORT_URL, {"file": SimpleUploadedFile(name, data)})


def preview(client, data: bytes, name="stations.csv") -> StationImport:
    response = upload(client, data, name)
    assert response.status_code == 302, response.content.decode()[:2000]
    return StationImport.objects.get(pk=response.url.rstrip("/").split("/")[-1])


def confirm(client, station_import: StationImport, fingerprint=None):
    return client.post(
        f"{IMPORT_URL}{station_import.pk}/confirm/",
        {"fingerprint": station_import.fingerprint if fingerprint is None else fingerprint},
    )


def make_station(**overrides) -> FuelStation:
    values = {
        "opis_id": 100,
        "name": "EXISTING STOP",
        "address": "I-35, EXIT 1",
        "city": "Tomah",
        "state": "WI",
        "rack_id": 420,
        "retail_price": Decimal("3.45"),
        "max_retail_price": Decimal("3.45"),
        "price_observations": 1,
        "latitude": 43.988465,
        "longitude": -90.501813,
        "geolocation_source": GeolocationSource.CENSUS_GAZETTEER_PLACES_2010,
        "geolocation_accuracy": GeolocationAccuracy.CITY,
        "location_uncertainty_miles": 1.6,
    }
    values.update(overrides)
    return FuelStation.objects.create(**values)


NEW_ROW = ["500", "NEW STOP", "I-8, EXIT 119", "Gila Bend", "AZ", "930", "3.899"]


# --- 1-3: access and permissions -------------------------------------------------------------


def test_admin_import_page_renders_for_authorized_staff(admin_client):
    response = admin_client.get(IMPORT_URL)
    html = response.content.decode()
    assert response.status_code == 200
    assert "Import Fuel Stations" in html
    assert 'enctype="multipart/form-data"' in html
    assert 'name="csrfmiddlewaretoken"' in html
    assert "Download Excel template" in html
    assert "no-store" in response.headers["Cache-Control"]


def test_changelist_shows_import_button_next_to_add(admin_client):
    html = admin_client.get("/admin/fuel/fuelstation/").content.decode()
    assert 'href="/admin/fuel/fuelstation/import/" class="importlink">Import Fuel Stations' in html
    assert html.index("addlink") < html.index("importlink")


@pytest.mark.parametrize(
    ("method", "path"),
    [
        ("get", IMPORT_URL),
        ("post", IMPORT_URL),
        ("get", f"{IMPORT_URL}template.xlsx"),
        ("get", f"{IMPORT_URL}00000000-0000-0000-0000-000000000000/"),
        ("post", f"{IMPORT_URL}00000000-0000-0000-0000-000000000000/confirm/"),
        ("get", f"{IMPORT_URL}00000000-0000-0000-0000-000000000000/errors.csv"),
    ],
)
def test_anonymous_requests_are_redirected_to_login(client, method, path):
    response = getattr(client, method)(path)
    assert response.status_code == 302
    assert response.url.startswith("/admin/login/?next=")
    assert not StationImport.objects.exists()


def test_staff_without_add_or_change_permission_is_forbidden(client, django_user_model):
    user, _ = make_user(django_user_model, "view_fuelstation", name="viewer")
    client.force_login(user)
    assert client.get(IMPORT_URL).status_code == 403
    assert upload(client, csv_bytes([NEW_ROW])).status_code == 403
    assert client.get(f"{IMPORT_URL}template.csv").status_code == 403
    html = client.get("/admin/fuel/fuelstation/").content.decode()
    assert "Import Fuel Stations" not in html
    assert not StationImport.objects.exists()


def test_add_only_staff_cannot_confirm_updates(client, django_user_model):
    make_station()
    user, _ = make_user(django_user_model, "add_fuelstation", name="adder")
    client.force_login(user)
    pending = preview(client, csv_bytes([["100", "", "", "", "", "", "3.29"]]))
    html = client.get(f"{IMPORT_URL}{pending.pk}/").content.decode()
    assert "Confirmation unavailable" in html and "change fuel stations" in html
    response = confirm(client, pending)
    assert response.status_code == 302
    assert FuelStation.objects.get().retail_price == Decimal("3.45")
    pending.refresh_from_db()
    assert pending.status == StationImport.Status.PREVIEWED


def test_change_only_staff_cannot_confirm_new_stations(client, django_user_model):
    user, _ = make_user(django_user_model, "change_fuelstation", name="changer")
    client.force_login(user)
    pending = preview(client, csv_bytes([NEW_ROW]))
    confirm(client, pending)
    assert not FuelStation.objects.exists()


def test_imports_are_private_to_their_uploader(admin_client, django_user_model):
    pending = preview(admin_client, csv_bytes([NEW_ROW]))
    other_user, _ = make_user(
        django_user_model, "add_fuelstation", "change_fuelstation", name="other"
    )
    other = Client()
    other.force_login(other_user)
    assert other.get(f"{IMPORT_URL}{pending.pk}/").status_code == 404
    assert other.get(f"{IMPORT_URL}{pending.pk}/errors.csv").status_code == 404
    assert confirm(other, pending).status_code == 404
    assert not FuelStation.objects.exists()


def test_confirmation_requires_post_and_csrf(admin_user):
    client = Client(enforce_csrf_checks=True)
    client.force_login(admin_user[0])
    token_page = client.get(IMPORT_URL)
    token = token_page.cookies["csrftoken"].value
    response = client.post(
        IMPORT_URL,
        {"file": SimpleUploadedFile("s.csv", csv_bytes([NEW_ROW])), "csrfmiddlewaretoken": token},
    )
    pending = StationImport.objects.get()
    assert response.status_code == 302
    assert client.get(f"{IMPORT_URL}{pending.pk}/confirm/").status_code == 405
    no_token = client.post(
        f"{IMPORT_URL}{pending.pk}/confirm/", {"fingerprint": pending.fingerprint}
    )
    assert no_token.status_code == 403
    assert not FuelStation.objects.exists()
    response = client.post(
        f"{IMPORT_URL}{pending.pk}/confirm/",
        {"fingerprint": pending.fingerprint, "csrfmiddlewaretoken": token},
    )
    assert response.status_code == 302
    assert FuelStation.objects.count() == 1


def test_upload_without_csrf_token_is_rejected(admin_user):
    client = Client(enforce_csrf_checks=True)
    client.force_login(admin_user[0])
    response = upload(client, csv_bytes([NEW_ROW]))
    assert response.status_code == 403
    assert not StationImport.objects.exists()


# --- 4-7: file formats and limits ------------------------------------------------------------


def test_csv_upload_previews_and_imports_new_station(admin_client):
    pending = preview(admin_client, csv_bytes([NEW_ROW]))
    assert pending.file_format == "csv"
    assert pending.summary["new"] == 1 and pending.summary["total_rows"] == 1
    assert not FuelStation.objects.exists()

    response = confirm(admin_client, pending)
    assert response.status_code == 302
    station = FuelStation.objects.get(opis_id=500)
    assert (station.name, station.city, station.state, station.rack_id) == (
        "NEW STOP",
        "Gila Bend",
        "AZ",
        930,
    )
    assert station.retail_price == station.max_retail_price == Decimal("3.899")
    assert station.price_observations == 1
    results = admin_client.get(response.url).content.decode()
    assert "Import applied" in results and "Rows processed" in results


def test_utf8_bom_and_case_insensitive_headers_are_accepted(admin_client):
    data = "﻿" + csv_bytes([NEW_ROW], header=[h.upper() for h in HEADER]).decode()
    pending = preview(admin_client, data.encode())
    assert pending.summary["new"] == 1


def test_xlsx_upload_with_numeric_cells(admin_client):
    data = xlsx_bytes([[500, "NEW STOP", "I-8, EXIT 119", "Gila Bend", "AZ", 930, 3.899]])
    pending = preview(admin_client, data, "stations.xlsx")
    assert pending.file_format == "xlsx"
    assert pending.summary["new"] == 1
    confirm(admin_client, pending)
    assert FuelStation.objects.get(opis_id=500).retail_price == Decimal("3.899")


def test_xlsx_reads_the_stations_sheet_and_ignores_examples(admin_client):
    data = xlsx_bytes(
        [[500, "NEW STOP", "", "Gila Bend", "AZ", 930, 3.899]],
        sheet="Stations",
        extra_sheets=[("Example", [HEADER, [999, "EXAMPLE", "", "Dallas", "TX", 1, 3.0]])],
    )
    pending = preview(admin_client, data, "stations.xlsx")
    assert [r["opis_id"] for r in pending.rows] == [500]


def test_xlsx_formula_date_and_error_cells_are_rejected_not_evaluated(admin_client):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Stations"
    sheet.append(HEADER)
    sheet.append([500, "FORMULA STOP", "", "Gila Bend", "AZ", 930, "=1+2.5"])
    sheet.append([501, "DATE STOP", "", "Gila Bend", "AZ", 930, None])
    sheet["G3"].value = __import__("datetime").date(2026, 1, 1)
    sheet.append([502, "GOOD STOP", "", "Gila Bend", "AZ", 930, 3.5])
    buffer = io.BytesIO()
    workbook.save(buffer)
    pending = preview(admin_client, buffer.getvalue(), "s.xlsx")
    reasons = {issue["opis_id"]: issue for issue in pending.issues}
    assert reasons["500"]["reason"] == "unsupported_value"
    assert "formula" in reasons["500"]["detail"]
    assert "date" in reasons["501"]["detail"]
    assert pending.summary["new"] == 1


@pytest.mark.parametrize(
    ("name", "data", "message"),
    [
        ("stations.txt", b"a,b\n1,2\n", "Unsupported file type"),
        ("stations.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 100, "Legacy .xls"),
        ("stations.xlsm", b"PK\x03\x04" + b"0" * 100, "Macro-enabled"),
        ("stations.csv", b"PK\x03\x04" + b"0" * 100, "binary data"),
        ("stations.csv", b"\xff\xfeO\x00P\x00", "binary data"),
        ("stations.csv", "OPIS Truckstop ID,City\n1,Montr\xe9al\n".encode("latin-1"), "UTF-8"),
        ("stations.xlsx", csv_bytes([NEW_ROW]), "not an Excel workbook"),
        ("stations.xlsx", b"PK\x03\x04garbage", "damaged"),
        ("stations.csv", b"", "empty"),
    ],
)
def test_unsupported_or_disguised_files_are_rejected(admin_client, name, data, message):
    response = upload(admin_client, data, name)
    assert response.status_code == 400
    assert message in response.content.decode()
    assert not StationImport.objects.exists()


def test_macro_enabled_content_in_xlsx_is_rejected(admin_client):
    original = xlsx_bytes([[500, "X", "", "Gila Bend", "AZ", 1, 3.5]])
    buffer = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(original)) as source,
        zipfile.ZipFile(buffer, "w") as target,
    ):
        for entry in source.infolist():
            target.writestr(entry, source.read(entry.filename))
        target.writestr("xl/vbaProject.bin", b"\x00" * 32)
    response = upload(admin_client, buffer.getvalue(), "stations.xlsx")
    assert response.status_code == 400
    assert "Macro-enabled" in response.content.decode()


def test_oversized_file_is_rejected_before_parsing(admin_client):
    data = csv_bytes([NEW_ROW]) + b"#" * station_import.MAX_UPLOAD_BYTES
    response = upload(admin_client, data)
    assert response.status_code == 400
    assert "larger than the 5 MB limit" in response.content.decode()
    assert not StationImport.objects.exists()


def test_row_and_expansion_limits(admin_client, monkeypatch):
    monkeypatch.setattr(station_import, "MAX_DATA_ROWS", 2)
    response = upload(admin_client, csv_bytes([NEW_ROW] * 3))
    assert "more than 2 data rows" in response.content.decode()
    monkeypatch.setattr(station_import, "MAX_XLSX_UNCOMPRESSED_BYTES", 1000)
    response = upload(admin_client, xlsx_bytes([NEW_ROW]), "s.xlsx")
    assert "expands to more data" in response.content.decode()
    assert not StationImport.objects.exists()


@pytest.mark.parametrize(
    ("header", "message"),
    [
        (["Truckstop Name", "City"], "required column “OPIS Truckstop ID” was not found"),
        (["OPIS Truckstop ID;Truckstop Name"], "semicolon-separated"),
        (["OPIS Truckstop ID"], "only has the OPIS Truckstop ID column"),
        (["OPIS Truckstop ID", "City", "city"], "appears more than once"),
    ],
)
def test_header_problems_reject_the_whole_file(admin_client, header, message):
    response = upload(admin_client, csv_bytes([], header=header))
    assert response.status_code == 400
    assert message in response.content.decode()


def test_downloadable_templates_match_the_import_columns(admin_client):
    response = admin_client.get(f"{IMPORT_URL}template.xlsx")
    assert response.status_code == 200
    assert "attachment" in response["Content-Disposition"]
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Stations", "Example", "Field guide"]
    assert [c.value for c in workbook["Stations"][1]] == HEADER
    assert workbook["Stations"].max_row == 1
    assert "never imported" in workbook["Example"]["A1"].value

    response = admin_client.get(f"{IMPORT_URL}template.csv")
    assert next(csv.reader(io.StringIO(response.content.decode()))) == HEADER
    # The unchanged template imports nothing.
    pending = preview(
        admin_client, admin_client.get(f"{IMPORT_URL}template.xlsx").content, "t.xlsx"
    )
    assert pending.summary["total_rows"] == 0
    assert "Nothing to import" in admin_client.get(f"{IMPORT_URL}{pending.pk}/").content.decode()


# --- 8-12: insert / update / skip / idempotency ----------------------------------------------


def test_existing_station_price_update_shows_current_and_new_values(admin_client):
    station = make_station()
    pending = preview(admin_client, csv_bytes([["100", "", "", "", "", "", "3.29"]]))
    assert pending.summary["updates"] == 1 and pending.summary["new"] == 0
    html = admin_client.get(f"{IMPORT_URL}{pending.pk}/").content.decode()
    assert "Retail Price" in html and "$3.45" in html and "$3.29" in html

    confirm(admin_client, pending)
    station.refresh_from_db()
    assert station.retail_price == station.max_retail_price == Decimal("3.29")
    assert station.name == "EXISTING STOP" and station.city == "Tomah"
    assert station.latitude == pytest.approx(43.988465)  # price change keeps coordinates
    assert FuelStation.objects.count() == 1


def test_existing_station_location_update_resets_and_rederives_coordinates(admin_client):
    station = make_station()
    row = ["100", "", "I-8, EXIT 119", "Gila Bend", "AZ", "", ""]
    pending = preview(admin_client, csv_bytes([row]))
    html = admin_client.get(f"{IMPORT_URL}{pending.pk}/").content.decode()
    assert "Tomah" in html and "Gila Bend" in html
    assert "current coordinates are removed and re-derived" in html
    confirm(admin_client, pending)
    station.refresh_from_db()
    assert (station.city, station.state, station.address) == ("Gila Bend", "AZ", "I-8, EXIT 119")
    assert station.geolocation_accuracy == GeolocationAccuracy.CITY
    assert station.latitude == pytest.approx(32.948625)  # the new town, not Tomah
    assert station.retail_price == Decimal("3.45")


def test_unchanged_stations_are_skipped_without_writes(admin_client):
    station = make_station()
    before = station.updated_at
    row = ["100", "EXISTING  STOP", "I-35, EXIT 1", "Tomah", "wi", "420", "3.4500"]
    pending = preview(admin_client, csv_bytes([row]))
    assert pending.summary["unchanged"] == 1
    assert pending.summary["updates"] == 0 and pending.summary["new"] == 0
    html = admin_client.get(f"{IMPORT_URL}{pending.pk}/").content.decode()
    assert "Nothing to import" in html and "Confirm import" not in html
    station.refresh_from_db()
    assert station.updated_at == before


def test_repeated_import_is_idempotent(admin_client):
    make_station()
    rows = [["100", "", "", "", "", "", "3.29"], NEW_ROW]
    first = preview(admin_client, csv_bytes(rows))
    confirm(admin_client, first)
    snapshot = list(FuelStation.objects.values_list("opis_id", "updated_at", "retail_price"))

    second = preview(admin_client, csv_bytes(rows))
    assert (second.summary["new"], second.summary["updates"], second.summary["unchanged"]) == (
        0,
        0,
        2,
    )
    confirm(admin_client, second)  # nothing to apply; harmless
    assert list(FuelStation.objects.values_list("opis_id", "updated_at", "retail_price")) == (
        snapshot
    )
    assert FuelStation.objects.count() == 2


# --- 13-17: validation, duplicates, partial columns, blanks ----------------------------------


@pytest.mark.parametrize(
    ("opis_id", "reason"),
    [("", "missing_field"), ("abc", "invalid_id"), ("-5", "invalid_id"), ("0", "invalid_id")],
)
def test_missing_or_invalid_identifiers_are_rejected(admin_client, opis_id, reason):
    pending = preview(admin_client, csv_bytes([[opis_id, *NEW_ROW[1:]], NEW_ROW]))
    assert [i["reason"] for i in pending.issues] == [reason]
    assert pending.issues[0]["line"] == 2
    assert pending.summary["new"] == 1


@pytest.mark.parametrize("price", ["abc", "0", "0.49", "20.01", "3.123456789", "NaN", "$3.29"])
def test_invalid_prices_are_rejected(admin_client, price):
    make_station()
    pending = preview(admin_client, csv_bytes([["100", "", "", "", "", "", price]]))
    assert [i["reason"] for i in pending.issues] == ["invalid_price"]
    assert pending.summary["updates"] == 0


def test_other_invalid_values_are_rejected_with_reasons(admin_client):
    rows = [
        ["501", "STOP", "", "Gila Bend", "ZZ", "930", "3.5"],
        ["502", "STOP", "", "Calgary", "AB", "930", "3.5"],
        ["503", "STOP", "", "Gila Bend", "AZ", "x", "3.5"],
        ["504", "S" * 256, "", "Gila Bend", "AZ", "930", "3.5"],
        ["505", "STOP", "", "Gila Bend", "AZ", "930"],
    ]
    pending = preview(admin_client, csv_bytes(rows))
    assert [i["reason"] for i in pending.issues] == [
        "invalid_state",
        "non_us_station",
        "invalid_rack_id",
        "value_too_long",
        "malformed_row",
    ]
    assert pending.summary["new"] == 0


def test_duplicate_ids_merge_or_reject_without_duplicate_stations(admin_client):
    rows = [
        ["600", "PILOT", "I-8", "Gila Bend", "AZ", "930", "3.899"],
        ["600", "PILOT TRAVEL CENTER", "i-8", "Gila Bend", "AZ", "930", "3.799"],
        ["600", "PILOT TRAVEL CENTER", "I-8", "gila bend", "AZ", "930", "3.999"],
        ["601", "LOVES", "I-65", "Gila Bend", "AZ", "250", "3.4"],
        ["601", "LOVES", "I-70", "Gila Bend", "AZ", "250", "3.5"],
    ]
    pending = preview(admin_client, csv_bytes(rows))
    assert pending.summary["duplicate_rows"] == 4
    assert pending.summary["duplicate_rows_merged"] == 2
    assert pending.summary["conflicting_duplicate_rows"] == 2
    assert {i["opis_id"] for i in pending.issues} == {"601"}
    confirm(admin_client, pending)
    station = FuelStation.objects.get()
    assert station.opis_id == 600 and station.name == "PILOT TRAVEL CENTER"
    assert (station.retail_price, station.max_retail_price, station.price_observations) == (
        Decimal("3.799"),
        Decimal("3.999"),
        3,
    )


def test_partial_columns_update_only_supplied_fields(admin_client):
    station = make_station()
    data = csv_bytes(
        [["100", "3.10"], ["700", "3.20"]], header=["OPIS Truckstop ID", "Retail Price"]
    )
    pending = preview(admin_client, data)
    assert pending.summary["updates"] == 1
    assert [i["reason"] for i in pending.issues] == ["incomplete_new_station"]
    assert "Truckstop Name" in pending.issues[0]["detail"]
    confirm(admin_client, pending)
    station.refresh_from_db()
    assert station.retail_price == Decimal("3.10")
    assert (station.name, station.address, station.city, station.rack_id) == (
        "EXISTING STOP",
        "I-35, EXIT 1",
        "Tomah",
        420,
    )
    assert not FuelStation.objects.filter(opis_id=700).exists()


def test_blank_cells_keep_stored_values(admin_client):
    station = make_station()
    rows = [
        ["100", "", "-", "", "", "", ""],  # all blank or placeholder: nothing changes
        ["800", "NEW STOP", "", "Gila Bend", "AZ", "930", ""],  # new station needs a price
    ]
    pending = preview(admin_client, csv_bytes(rows))
    assert pending.summary["unchanged"] == 1
    assert [i["reason"] for i in pending.issues] == ["incomplete_new_station"]
    assert "Retail Price" in pending.issues[0]["detail"]
    station.refresh_from_db()
    assert station.address == "I-35, EXIT 1"


def test_new_station_without_address_is_allowed(admin_client):
    pending = preview(admin_client, csv_bytes([["801", "S", "", "Gila Bend", "AZ", "9", "3.5"]]))
    confirm(admin_client, pending)
    assert FuelStation.objects.get().address == ""


def test_unknown_and_coordinate_columns_are_ignored_with_a_note(admin_client):
    header = [*HEADER, "Latitude", "Longitude"]
    pending = preview(admin_client, csv_bytes([[*NEW_ROW, "1.0", "2.0"]], header=header))
    assert any("Latitude" in note for note in pending.summary["notes"])
    assert any("Coordinates are never imported" in note for note in pending.summary["notes"])
    confirm(admin_client, pending)
    station = FuelStation.objects.get()
    assert station.latitude == pytest.approx(32.948625)  # from the gazetteer, not the file


# --- 18-19: geolocation -----------------------------------------------------------------------


def test_geolocation_metadata_is_preserved_on_non_location_updates(admin_client):
    address_level = make_station(
        opis_id=101,
        geolocation_accuracy=GeolocationAccuracy.ADDRESS,
        latitude=43.9,
        longitude=-90.4,
        location_uncertainty_miles=None,
    )
    city_level = make_station(opis_id=102)
    rows = [["101", "RENAMED", "", "", "", "421", "3.1"], ["102", "", "", "", "", "", "3.2"]]
    confirm(admin_client, preview(admin_client, csv_bytes(rows)))
    address_level.refresh_from_db()
    city_level.refresh_from_db()
    assert (address_level.latitude, address_level.longitude) == (43.9, -90.4)
    assert address_level.geolocation_accuracy == GeolocationAccuracy.ADDRESS
    assert address_level.name == "RENAMED" and address_level.rack_id == 421
    assert city_level.geolocation_source == GeolocationSource.CENSUS_GAZETTEER_PLACES_2010
    assert city_level.location_uncertainty_miles == 1.6


def test_relocated_station_never_keeps_stale_coordinates(admin_client):
    verified = make_station(
        opis_id=103,
        geolocation_accuracy=GeolocationAccuracy.ADDRESS,
        latitude=43.9,
        longitude=-90.4,
        location_uncertainty_miles=None,
    )
    unknown_town = make_station(opis_id=104)
    rows = [
        ["103", "", "NEW ADDRESS 5", "", "", "", ""],
        ["104", "", "", "Nowhereville", "NE", "", ""],
    ]
    confirm(admin_client, preview(admin_client, csv_bytes(rows)))
    verified.refresh_from_db()
    unknown_town.refresh_from_db()
    # Address moved: the verified point is gone; only the town approximation is re-derived.
    assert verified.geolocation_accuracy == GeolocationAccuracy.CITY
    assert verified.latitude == pytest.approx(43.988465)
    # Town not in the gazetteer: unresolved, no coordinates, excluded from routing.
    assert unknown_town.geolocation_accuracy == GeolocationAccuracy.UNRESOLVED
    assert unknown_town.latitude is None and unknown_town.geolocation_source == ""
    assert not routable_stations().filter(pk=unknown_town.pk).exists()


def test_new_stations_become_routable_and_change_the_station_data_version(admin_client):
    before = _station_data_version()
    rows = [NEW_ROW, ["900", "NOWHERE STOP", "", "Nowhereville", "NE", "1", "3.0"]]
    confirm(admin_client, preview(admin_client, csv_bytes(rows)))
    assert routable_stations().filter(opis_id=500).exists()
    nowhere = FuelStation.objects.get(opis_id=900)
    assert nowhere.geolocation_accuracy == GeolocationAccuracy.UNRESOLVED
    assert nowhere.latitude is None
    assert _station_data_version() != before
    result = StationImport.objects.get().result
    assert (result["geolocated"], result["left_unresolved"]) == (1, 1)


def test_missing_gazetteer_leaves_new_stations_unresolved(admin_client, settings, tmp_path):
    settings.GAZETTEER_PLACES_PATH = tmp_path / "missing.txt"
    confirm(admin_client, preview(admin_client, csv_bytes([NEW_ROW])))
    station = FuelStation.objects.get()
    assert station.geolocation_accuracy == GeolocationAccuracy.UNRESOLVED
    assert StationImport.objects.get().result["geolocation"] == "skipped"


# --- 20-23: preview, transactions, double submit, concurrency --------------------------------


def test_preview_does_not_modify_station_data(admin_client):
    make_station()
    before = list(FuelStation.objects.values())
    rows = [["100", "", "", "Gila Bend", "AZ", "", "3.10"], NEW_ROW]
    with CaptureQueriesContext(connection) as queries:
        pending = preview(admin_client, csv_bytes(rows))
    station_writes = [
        q["sql"]
        for q in queries.captured_queries
        if "fuel_fuelstation" in q["sql"] and q["sql"].split()[0] in {"INSERT", "UPDATE", "DELETE"}
    ]
    assert station_writes == []
    assert list(FuelStation.objects.values()) == before
    assert pending.status == StationImport.Status.PREVIEWED
    station_lookups = [q for q in queries.captured_queries if 'FROM "fuel_fuelstation"' in q["sql"]]
    assert len(station_lookups) == 1  # one batched match, not one query per row


def test_confirmation_is_all_or_nothing(admin_client, admin_user, monkeypatch):
    make_station()
    pending = preview(admin_client, csv_bytes([["100", "", "", "", "", "", "3.10"], NEW_ROW]))

    def explode(*args, **kwargs):
        raise RuntimeError("simulated failure after writes")

    monkeypatch.setattr(station_import, "geolocate_stations", explode)
    with pytest.raises(RuntimeError):
        station_import.confirm_import(pending.pk, admin_user[0], pending.fingerprint)
    assert FuelStation.objects.count() == 1
    assert FuelStation.objects.get().retail_price == Decimal("3.45")
    pending.refresh_from_db()
    assert pending.status == StationImport.Status.PREVIEWED  # claim rolled back; can retry

    monkeypatch.undo()
    station_import.confirm_import(pending.pk, admin_user[0], pending.fingerprint)
    assert FuelStation.objects.count() == 2


def test_double_submission_applies_once(admin_client):
    pending = preview(admin_client, csv_bytes([NEW_ROW]))
    first = confirm(admin_client, pending)
    second = confirm(admin_client, pending)
    assert first.status_code == second.status_code == 302
    assert FuelStation.objects.count() == 1
    page = admin_client.get(second.url).content.decode()
    assert "already confirmed" in page
    pending.refresh_from_db()
    assert pending.status == StationImport.Status.APPLIED
    assert pending.rows == []


def test_tampered_fingerprint_is_refused(admin_client):
    pending = preview(admin_client, csv_bytes([NEW_ROW]))
    confirm(admin_client, pending, fingerprint="0" * 64)
    assert not FuelStation.objects.exists()
    pending.refresh_from_db()
    assert pending.status == StationImport.Status.PREVIEWED


def test_concurrent_edit_after_preview_makes_it_stale(admin_client):
    station = make_station()
    pending = preview(admin_client, csv_bytes([["100", "", "", "", "", "", "3.10"], NEW_ROW]))
    FuelStation.objects.filter(pk=station.pk).update(name="EDITED ELSEWHERE")  # someone else
    FuelStation.objects.filter(pk=station.pk).update(updated_at=station.updated_at + timedelta(1))
    response = confirm(admin_client, pending)
    page = admin_client.get(response.url).content.decode()
    assert "Station data changed after this preview" in page
    assert "Preview out of date" in page
    assert not FuelStation.objects.filter(opis_id=500).exists()
    station.refresh_from_db()
    assert station.retail_price == Decimal("3.45") and station.name == "EDITED ELSEWHERE"
    pending.refresh_from_db()
    assert pending.status == StationImport.Status.STALE


def test_two_previews_of_the_same_file_cannot_both_apply(admin_client):
    first = preview(admin_client, csv_bytes([NEW_ROW]))
    second = preview(admin_client, csv_bytes([NEW_ROW]))
    confirm(admin_client, first)
    confirm(admin_client, second)
    assert FuelStation.objects.count() == 1
    second.refresh_from_db()
    assert second.status == StationImport.Status.STALE


def test_expired_preview_cannot_be_confirmed(admin_client):
    pending = preview(admin_client, csv_bytes([NEW_ROW]))
    StationImport.objects.filter(pk=pending.pk).update(
        created_at=pending.created_at - station_import.PREVIEW_TTL - timedelta(seconds=1)
    )
    assert "Preview expired" in admin_client.get(f"{IMPORT_URL}{pending.pk}/").content.decode()
    confirm(admin_client, pending)
    assert not FuelStation.objects.exists()
    pending.refresh_from_db()
    assert pending.status == StationImport.Status.EXPIRED


# --- 24: error report -------------------------------------------------------------------------


def test_error_report_lists_rejected_rows_and_neutralizes_formulas(admin_client):
    rows = [["=HYPERLINK(1)", *NEW_ROW[1:]], ["", *NEW_ROW[1:]], NEW_ROW]
    pending = preview(admin_client, csv_bytes(rows))
    html = admin_client.get(f"{IMPORT_URL}{pending.pk}/").content.decode()
    assert "Invalid rows" in html and "Download error report (CSV)" in html
    response = admin_client.get(f"{IMPORT_URL}{pending.pk}/errors.csv")
    assert response.status_code == 200
    assert response["Content-Type"].startswith("text/csv")
    assert "attachment" in response["Content-Disposition"]
    report = list(csv.reader(io.StringIO(response.content.decode())))
    assert report[0] == ["Row", "OPIS Truckstop ID", "Reason", "Code", "Detail"]
    assert [row[0] for row in report[1:]] == ["2", "3"]
    assert report[1][1] == "'=HYPERLINK(1)"  # formula injection neutralized
    assert report[1][3] == "invalid_id" and report[2][3] == "missing_field"
    assert "Traceback" not in response.content.decode()


# --- 25: compatibility with the existing CLI importer ----------------------------------------


def test_cli_import_followed_by_admin_import_of_same_file_changes_nothing(admin_client, sample_csv):
    call_command("import_fuel_prices", file=sample_csv, skip_geolocation=True, stdout=io.StringIO())
    snapshot = list(FuelStation.objects.order_by("opis_id").values())
    pending = preview(admin_client, sample_csv.read_bytes())
    assert (pending.summary["new"], pending.summary["updates"]) == (0, 0)
    assert pending.summary["unchanged"] == FuelStation.objects.count() == 6
    assert pending.summary["duplicate_rows_merged"] == 2
    assert pending.summary["invalid_by_reason"] == {"non_us_station": 1}
    assert list(FuelStation.objects.order_by("opis_id").values()) == snapshot


def test_admin_import_then_cli_import_of_same_file_reports_unchanged(admin_client, sample_csv):
    confirm(admin_client, preview(admin_client, sample_csv.read_bytes()))
    assert FuelStation.objects.count() == 6
    out = io.StringIO()
    call_command("import_fuel_prices", file=sample_csv, skip_geolocation=True, stdout=out)
    assert "Unchanged:               6" in out.getvalue()
    assert "Inserted:                0" in out.getvalue()
